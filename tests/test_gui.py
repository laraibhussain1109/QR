import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import time
import cv2
from PyQt6.QtWidgets import QApplication
from openpyxl import load_workbook

from app import MainWindow
from scanner import make_sample


def pump(application, condition, timeout=8):
    deadline = time.monotonic() + timeout
    while not condition() and time.monotonic() < deadline:
        application.processEvents()
        time.sleep(0.01)
    assert condition()


def test_camera_worker_preview_overlay_excel_and_stop(tmp_path, monkeypatch):
    application = QApplication.instance() or QApplication([])
    sample = tmp_path / "qr.png"
    make_sample(sample)
    frame = cv2.imread(str(sample))
    captures = []
    class FakeCapture:
        def __init__(self, *args):
            self.released = False
            captures.append(self)
        def isOpened(self):
            return True
        def read(self):
            return True, frame.copy()
        def release(self):
            self.released = True
    monkeypatch.setattr("app.cv2.VideoCapture", FakeCapture)
    output = tmp_path / "parts.xlsx"
    window = MainWindow(output=str(output))
    window.show()
    try:
        for _ in range(2):
            window.start_camera()
            pump(application, lambda: "Steel Bracket" in window.details.text()
                 and ("Saved" in window.status.text() or "duplicate" in window.status.text()))
            assert "B-12" in window.details.text()
            assert window.view.pixmap() is not None
            window.stop_camera()
            pump(application, lambda: window.worker is None)
            window.details.setText("")
        assert all(cap.released for cap in captures)
        wb = load_workbook(output)
        assert wb["manufacturing"].max_row == 2
        wb.close()
    finally:
        window.stop_camera()
        pump(application, lambda: window.worker is None)
        window.close()


def test_missing_camera_reports_error_and_releases(tmp_path, monkeypatch):
    application = QApplication.instance() or QApplication([])
    class MissingCapture:
        released = False
        def isOpened(self):
            return False
        def release(self):
            self.released = True
    capture = MissingCapture()
    monkeypatch.setattr("app.cv2.VideoCapture", lambda *args: capture)
    window = MainWindow(output=str(tmp_path / "parts.xlsx"))
    window.start_camera()
    pump(application, lambda: window.worker is None)
    assert "Camera unavailable" in window.status.text()
    assert window.start.isEnabled()
    assert capture.released
    window.close()
