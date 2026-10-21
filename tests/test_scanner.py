import cv2
import numpy as np
from openpyxl import load_workbook
import pytest

from parts import ExcelStore, Part
from scanner import Scanner, make_sample


def test_real_qr_to_excel(tmp_path):
    image = tmp_path / "label.png"
    make_sample(image)
    detections = Scanner().detect(cv2.imread(str(image)))
    assert len(detections) == 1
    assert detections[0].part == Part("Steel Bracket", "P-001", "B-12")
    path = tmp_path / "scans.xlsx"
    assert ExcelStore(path).add(detections[0].part).added
    assert not ExcelStore(path).add(detections[0].part).added
    wb = load_workbook(path)
    assert wb["manufacturing"].max_row == 2
    wb.close()


def test_blank_frame():
    assert Scanner().detect(np.full((400, 400, 3), 255, dtype=np.uint8)) == []


@pytest.mark.parametrize("payload,expected", [
    ("nut-0124-3", Part("nut", "0124", "3")),
    ("Part Name: Nut\nPart ID: 0124\nBin Number: 3", Part("Nut", "0124", "3")),
])
def test_actual_qr_decoding_for_additional_formats(tmp_path, payload, expected):
    qr = cv2.QRCodeEncoder_create().encode(payload)
    qr = cv2.copyMakeBorder(qr, 4, 4, 4, 4, cv2.BORDER_CONSTANT, value=255)
    qr = cv2.resize(qr, None, fx=10, fy=10, interpolation=cv2.INTER_NEAREST)
    detections = Scanner().detect(cv2.cvtColor(qr, cv2.COLOR_GRAY2BGR))
    assert len(detections) == 1
    assert detections[0].payload == payload
    assert detections[0].part == expected
    path = tmp_path / "scanned.xlsx"
    assert ExcelStore(path).add(detections[0].part).added
    assert not ExcelStore(path).add(detections[0].part).added
    wb = load_workbook(path)
    assert wb["manufacturing"].max_row == 2
    assert tuple(c.value for c in wb["manufacturing"][2])[:3] == (
        expected.name, expected.part_id, expected.bin_number)
    wb.close()
