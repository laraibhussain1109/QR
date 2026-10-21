import json
from concurrent.futures import ThreadPoolExecutor

from openpyxl import Workbook, load_workbook
import pytest

from parts import ExcelStore, HEADERS, Part, parse_payload


def test_formats_and_leading_zeroes():
    expected = Part("Steel Bracket", "001", "B-12")
    assert parse_payload(json.dumps(dict(part_name=expected.name, part_id="001", bin_number="B-12"))) == expected
    assert parse_payload(" Steel Bracket ,001,B-12") == expected
    assert parse_payload('"Bracket, large",001,B-12').name == "Bracket, large"


@pytest.mark.parametrize("payload", ["", "random text", "a,b", "a,b,c,d", "{bad}",
                                      '{"part_name":"A"}',
                                      '{"part_name":"A","part_id":1,"bin_number":"B"}',
                                      '"a\nb",id,bin', "a,,b", "a," + "b" * 257 + ",c"])
def test_invalid_labels(payload):
    with pytest.raises(ValueError):
        parse_payload(payload)


def test_restart_duplicate_and_conflict(tmp_path):
    path = tmp_path / "manufacturing.xlsx"
    part = Part("Bracket", "001", "B-12")
    ExcelStore(path).initialize()
    assert ExcelStore(path).add(part).added
    assert not ExcelStore(path).add(part).added
    result = ExcelStore(path).add(Part("Other", "001", "B-99"))
    assert not result.added and result.part == part
    assert ExcelStore(path).add(Part("Other", "002", "B-12")).added
    wb = load_workbook(path)
    ws = wb["manufacturing"]
    assert tuple(c.value for c in ws[1]) == HEADERS
    assert ws.max_row == 3
    assert ws["B2"].value == "001"
    assert ws["C2"].value == "B-12"
    assert ws["D2"].value and "+" in ws["D2"].value
    wb.close()


def test_formula_is_saved_as_text(tmp_path):
    path = tmp_path / "parts.xlsx"
    ExcelStore(path).add(Part('=HYPERLINK("bad")', "001", "=1+1"))
    wb = load_workbook(path)
    assert wb["manufacturing"]["A2"].data_type == "s"
    assert wb["manufacturing"]["C2"].data_type == "s"
    wb.close()


def test_existing_wrong_headers_not_overwritten(tmp_path):
    path = tmp_path / "parts.xlsx"
    wb = Workbook()
    wb.active.title = "manufacturing"
    wb.active.append(["Unexpected"])
    wb.save(path)
    before = path.read_bytes()
    with pytest.raises(ValueError, match="headers"):
        ExcelStore(path).add(Part("a", "b", "c"))
    assert path.read_bytes() == before


def test_failed_replace_keeps_workbook(tmp_path, monkeypatch):
    path = tmp_path / "parts.xlsx"
    store = ExcelStore(path)
    store.initialize()
    before = path.read_bytes()
    def denied(*args):
        raise PermissionError("Workbook open in Excel")
    monkeypatch.setattr("parts.os.replace", denied)
    with pytest.raises(PermissionError):
        store.add(Part("a", "b", "c"))
    assert path.read_bytes() == before
    assert not list(tmp_path.glob(".parts-*"))


def test_concurrent_writers_do_not_duplicate(tmp_path):
    path = tmp_path / "parts.xlsx"
    def write(_):
        return ExcelStore(path).add(Part("a", "b", "c")).added
    with ThreadPoolExecutor(max_workers=4) as pool:
        assert sum(pool.map(write, range(8))) == 1
    wb = load_workbook(path)
    assert wb["manufacturing"].max_row == 2
    wb.close()
