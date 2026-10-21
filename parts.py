"""QR payload validation and durable, duplicate-safe Excel storage."""
from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime
import io
import json
import os
from pathlib import Path
import re
import tempfile

from filelock import FileLock
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

HEADERS = ("Part Name", "Part ID", "BIN Number", "TimeStamp")
SHEET = "manufacturing"


@dataclass(frozen=True)
class Part:
    name: str
    part_id: str
    bin_number: str

    def __post_init__(self):
        for value in (self.name, self.part_id, self.bin_number):
            if not isinstance(value, str) or not value.strip():
                raise ValueError("Part Name, Part ID and BIN Number must be nonempty strings.")
            if len(value) > 256 or re.search(r"[\x00-\x1f\x7f]", value):
                raise ValueError("Part fields must be at most 256 characters without control characters.")
        object.__setattr__(self, "name", self.name.strip())
        object.__setattr__(self, "part_id", self.part_id.strip())
        object.__setattr__(self, "bin_number", self.bin_number.strip())


FIELD_ALIASES = {
    "partname": "name",
    "partid": "part_id",
    "partnumber": "part_id",
    "partno": "part_id",
    "binnumber": "bin_number",
    "binno": "bin_number",
    "bin": "bin_number",
}
LABEL = re.compile(
    r"(?<![\w])(?P<key>part[ _-]*(?:name|id|number|no\.?)|"
    r"bin(?:[ _-]*(?:number|no\.?))?)\s*[:=][ \t]*",
    re.IGNORECASE,
)
COMPACT_LABEL = re.compile(
    r"(?P<name>\w[\w .-]*)-(?P<part_id>[0-9]+)-(?P<bin_number>[A-Za-z0-9]+)"
)


def _part_from_fields(fields) -> Part:
    values = {}
    for key, value in fields:
        normalized = re.sub(r"[\s_.-]", "", key).casefold()
        field = FIELD_ALIASES.get(normalized)
        if field is None:
            continue
        if field in values:
            raise ValueError(f"QR contains more than one value for {key}; use one per field.")
        values[field] = value
    missing = [label for field, label in (("name", "Part Name"), ("part_id", "Part ID"),
                                         ("bin_number", "BIN Number")) if field not in values]
    if missing:
        raise ValueError("QR is missing fields: " + ", ".join(missing))
    return Part(**values)


def _json_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"Duplicate JSON field: {key}")
        result[key] = value
    return result


def parse_payload(payload: str) -> Part:
    """Read JSON, labelled text, name-numericID-bin, or one CSV row."""
    payload = payload.lstrip("\ufeff").strip()
    if len(payload) > 4096:
        raise ValueError("QR payload is too long.")
    if payload.startswith("{"):
        try:
            obj = json.loads(payload, object_pairs_hook=_json_object)
        except json.JSONDecodeError as exc:
            raise ValueError("QR contains invalid JSON; check quotes and punctuation.") from exc
        return _part_from_fields(obj.items())
    labels = list(LABEL.finditer(payload))
    # Quoted CSV is unambiguously positional, even if the part name contains a label.
    if labels and not payload.startswith('"'):
        fields = []
        for index, label in enumerate(labels):
            end = labels[index + 1].start() if index + 1 < len(labels) else len(payload)
            value = payload[label.end():end].strip().rstrip(",;|").strip()
            fields.append((label["key"], value))
        return _part_from_fields(fields)
    compact = COMPACT_LABEL.fullmatch(payload)
    if compact:
        return Part(**compact.groupdict())
    try:
        rows = list(csv.reader(io.StringIO(payload), strict=True))
    except csv.Error as exc:
        raise ValueError("Invalid CSV label.") from exc
    if len(rows) != 1 or len(rows[0]) != 3:
        raise ValueError("QR needs Part Name, Part ID and BIN Number as JSON, labelled text, CSV or name-ID-bin (nut-0124-3).")
    return Part(*rows[0])


@dataclass(frozen=True)
class SaveResult:
    added: bool
    part: Part


class ExcelStore:
    """One row per case-sensitive, trimmed Part ID, including across restarts.

    A sidecar lock serializes cooperating app instances. Atomic replacement keeps
    the original workbook intact if saving fails. Close Excel before scanning.
    """

    def __init__(self, path: str | Path):
        self.path = Path(path).expanduser().resolve()
        if self.path.suffix.lower() != ".xlsx":
            raise ValueError("Choose an .xlsx workbook.")
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.lock = FileLock(str(self.path) + ".lock", timeout=2)

    def _open(self):
        if self.path.exists():
            wb = load_workbook(self.path)
            if SHEET not in wb.sheetnames:
                wb.close()
                raise ValueError(f"Existing workbook needs a '{SHEET}' sheet with the required headers.")
            ws = wb[SHEET]
            if tuple(ws.cell(1, col).value for col in range(1, 5)) != HEADERS:
                wb.close()
                raise ValueError("Workbook headers must be: " + ", ".join(HEADERS))
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = SHEET
            ws.append(HEADERS)
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="245B78")
            for col, width in zip("ABCD", (32, 24, 22, 30)):
                ws.column_dimensions[col].width = width
        return wb, ws

    def _save(self, wb):
        fd, temporary = tempfile.mkstemp(prefix=".parts-", suffix=".xlsx", dir=self.path.parent)
        try:
            with os.fdopen(fd, "wb") as stream:
                wb.save(stream)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, self.path)
        finally:
            if os.path.exists(temporary):
                os.unlink(temporary)

    def initialize(self):
        with self.lock:
            wb, _ = self._open()
            try:
                if not self.path.exists():
                    self._save(wb)
            finally:
                wb.close()

    def add(self, part: Part) -> SaveResult:
        with self.lock:
            wb, ws = self._open()
            try:
                for row in ws.iter_rows(min_row=2, max_col=4, values_only=True):
                    if row[1] is not None and str(row[1]).strip() == part.part_id:
                        return SaveResult(False, Part(str(row[0]), str(row[1]), str(row[2])))
                ws.append((part.name, part.part_id, part.bin_number,
                           datetime.now().astimezone().isoformat(timespec="seconds")))
                # Explicit text cells preserve leading zeroes and prevent formula execution.
                for cell in ws[ws.max_row]:
                    cell.data_type = "s"
                    cell.number_format = "@"
                ws.auto_filter.ref = ws.dimensions
                self._save(wb)
                return SaveResult(True, part)
            finally:
                wb.close()
